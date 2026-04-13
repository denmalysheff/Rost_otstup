import streamlit as st
import pandas as pd

st.set_page_config(page_title="Рост отступлений", layout="wide")

# ==============================
# НОРМАЛИЗАЦИЯ
# ==============================
def normalize_columns(df):
    new_cols = {}

    for col in df.columns:
        c = col.strip().upper()

        if c in ["КМ", "KM"]:
            new_cols[col] = "KM"
        elif c in ["М", "M"]:
            new_cols[col] = "M"
        elif c in ["КОДНАПРВ"]:
            new_cols[col] = "KOD"
        elif c in ["ПУТЬ"]:
            new_cols[col] = "PATH"
        elif c in ["АМПЛИТУДА"]:
            new_cols[col] = "AMP"
        elif c in ["ДЛИНА"]:
            new_cols[col] = "LEN"
        elif c in ["СТЕПЕНЬ"]:
            new_cols[col] = "STEP"
        elif c in ["ОТСТУПЛЕНИЕ"]:
            new_cols[col] = "OTST"
        elif c in ["БАЛЛ"]:
            new_cols[col] = "BALL"
        elif c in ["ИС"]:
            new_cols[col] = "IS"
        elif c in ["СТРЕЛКА"]:
            new_cols[col] = "STR"
        elif c in ["МОСТ"]:
            new_cols[col] = "MOST"
        elif c in ["ГОД"]:
            new_cols[col] = "YEAR"
        elif c in ["МЕСЯЦ"]:
            new_cols[col] = "MONTH"
        elif c in ["ДЕНЬ"]:
            new_cols[col] = "DAY"
        else:
            new_cols[col] = col

    return df.rename(columns=new_cols)


def to_numeric(series):
    return pd.to_numeric(series.astype(str).str.replace(",", "."), errors="coerce")


def make_date(df):
    return pd.to_datetime(
        df["YEAR"].astype(str) + "-" +
        df["MONTH"].astype(str) + "-" +
        df["DAY"].astype(str),
        errors="coerce"
    )


# ==============================
# UI
# ==============================
st.title("📊 Рост отступлений")

file1 = st.file_uploader("Загрузите Файл 1", type=["xlsx"])
file2 = st.file_uploader("Загрузите Файл 2", type=["xlsx"])

tolerance = st.number_input("Допуск (м)", value=3)

# ==============================
# ОБРАБОТКА
# ==============================
if st.button("Запуск"):

    if not file1 or not file2:
        st.error("Загрузите оба файла")
        st.stop()

    try:
        df1 = pd.read_excel(file1, sheet_name="Отступления")
        df2 = pd.read_excel(file2, sheet_name="Отступления")

        df1 = normalize_columns(df1)
        df2 = normalize_columns(df2)

        df1["DATE"] = make_date(df1)
        df2["DATE"] = make_date(df2)

        # старый / новый
        if df1["DATE"].min() < df2["DATE"].min():
            old, new = df1, df2
        else:
            old, new = df2, df1

        # очистка
        for df in [old, new]:
            df["KOD"] = df["KOD"].astype(str).str.strip()
            df["PATH"] = df["PATH"].astype(str).str.strip()
            df["OTST"] = df["OTST"].astype(str).str.strip()

            for col in ["KM", "M", "AMP", "LEN", "BALL"]:
                if col in df.columns:
                    df[col] = to_numeric(df[col])

        old = old.dropna(subset=["KM", "M", "AMP"])
        new = new.dropna(subset=["KM", "M", "AMP"])

        # ==============================
        # MERGE
        # ==============================
        merged = pd.merge(
            old,
            new,
            on=["KOD", "PATH", "KM", "OTST"],
            suffixes=("_old", "_new")
        )

        merged["delta_m"] = abs(merged["M_new"] - merged["M_old"])
        merged = merged[merged["delta_m"] <= tolerance]

        merged = merged.sort_values("delta_m").drop_duplicates(
            subset=["KOD", "PATH", "KM", "M_old"],
            keep="first"
        )

        result = merged[merged["AMP_new"] > merged["AMP_old"]].copy()

        # ==============================
        # ФИНАЛ
        # ==============================
        def safe_col(df, name):
            return df[name] if name in df.columns else None

        df_result = pd.DataFrame({
            "Дата_старая": result["DATE_old"],
            "Дата_новая": result["DATE_new"],
            "Код направления": result["KOD"],
            "Путь": result["PATH"],
            "КМ": result["KM"],
            "М": result["M_old"],
            "Отступление": result["OTST"],
            "Степень": safe_col(result, "STEP_old"),
            "Амплитуда": result["AMP_old"],
            "Длина": safe_col(result, "LEN_old"),
            "Балл": safe_col(result, "BALL_old"),

            "М после": result["M_new"],

            "Степень после": safe_col(result, "STEP_new"),
            "Амплитуда после": result["AMP_new"],
            "Длина после": safe_col(result, "LEN_new"),
            "Балл после": safe_col(result, "BALL_new"),

            "ИС": safe_col(result, "IS_old"),
            "СТРЕЛКА": safe_col(result, "STR_old"),
            "МОСТ": safe_col(result, "MOST_old"),

            "Рост": result["AMP_new"] - result["AMP_old"]
        })

        st.success(f"Найдено записей: {len(df_result)}")

        st.dataframe(df_result, use_container_width=True)

        # ==============================
        # СКАЧИВАНИЕ
        # ==============================
        def to_excel(df):
            import io
            from openpyxl import Workbook
            from openpyxl.styles import PatternFill

            output = io.BytesIO()

            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Результат")

                ws = writer.book["Результат"]

                red = PatternFill(start_color="FF9999", fill_type="solid")
                yellow = PatternFill(start_color="FFFF99", fill_type="solid")

                growth_col = None
                for col in ws.iter_cols(1, ws.max_column):
                    if col[0].value == "Рост":
                        growth_col = col[0].column
                        break

                if growth_col:
                    for row in range(2, ws.max_row + 1):
                        val = ws.cell(row=row, column=growth_col).value

                        if val is None:
                            continue

                        if val > 5:
                            ws.cell(row=row, column=growth_col).fill = red
                        elif val > 2:
                            ws.cell(row=row, column=growth_col).fill = yellow

            output.seek(0)
            return output

        st.download_button(
            label="📥 Скачать Excel",
            data=to_excel(df_result),
            file_name="result.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        st.error(str(e))